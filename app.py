import streamlit as st
import pandas as pd
import io
import re
import csv
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

CSV_HEADER = [
    "Question", "Question Type",
    "Answer Option 1", "Explanation 1",
    "Answer Option 2", "Explanation 2",
    "Answer Option 3", "Explanation 3",
    "Answer Option 4", "Explanation 4",
    "Answer Option 5", "Explanation 5",
    "Answer Option 6", "Explanation 6",
    "Correct Answers", "Overall Explanation", "Domain"
]

def processar_questoes(texto, origem):
    questoes = []
    blocos = re.split(r'Question \d+', texto)

    for bloco in blocos:
        bloco = bloco.strip()
        if not bloco:
            continue

        linhas = bloco.split("\n")
        pergunta = ""
        opcoes = []
        respostas_corretas = []
        explicacao = ""
        captura_explicacao = False
        encontrou_true_false = False

        for i, linha in enumerate(linhas):
            linha = linha.strip()

            if i == 0 and linha.lower() == 'skipped':
                continue

            if i == 1 and linha != '':
                pergunta = linha
                continue

            if linha.lower() in [
                'choose the correct answer.',
                'there are two correct answers.',
                'there are three correct answers.',
                'there are four correct answers.',
                'there are five correct answers.'
            ]:
                continue

            if linha.lower().startswith('correct answer') or linha.lower().startswith('correct selection'):
                try:
                    resposta = linhas[i+1].strip()
                    if resposta:
                        respostas_corretas.append(resposta)
                except:
                    pass

            elif linha.lower().startswith('overall explanation'):
                captura_explicacao = True

            elif captura_explicacao:
                explicacao += linha + " "

            else:
                if linha and not linha.lower().startswith('note') and not linha.lower().startswith('skipped'):
                    if linha.lower() in ['true', 'false']:
                        encontrou_true_false = True
                    opcoes.append(linha)

        if encontrou_true_false and not opcoes:
            opcoes = ["True", "False"]

        if not pergunta.strip() or all(not alt.strip() for alt in opcoes):
            continue

        while len(opcoes) < 5:
            opcoes.append("")

        pergunta_formatada = pergunta
        if len(respostas_corretas) > 1:
            pergunta_formatada += f" ({len(respostas_corretas)} correct)"

        questoes.append({
            "Pergunta": pergunta_formatada,
            "Opção A": opcoes[0],
            "Opção B": opcoes[1],
            "Opção C": opcoes[2],
            "Opção D": opcoes[3],
            "Opção E": opcoes[4],
            "Resposta(s) Correta(s)": "; ".join(respostas_corretas),
            "Explicação": explicacao.strip(),
            "Origem": origem
        })

    return questoes

def gerar_xlsx(questoes, nome_arquivo):
    output = io.BytesIO()
    df_final = pd.DataFrame(questoes)
    df_final = df_final.sort_values(by="Pergunta").reset_index(drop=True)

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_final.to_excel(writer, index=False, sheet_name='Questões')
        worksheet = writer.sheets['Questões']
        header_fill = PatternFill(start_color="B7DEE8", end_color="B7DEE8", fill_type="solid")
        for col_num, _ in enumerate(df_final.columns, 1):
            col_letter = get_column_letter(col_num)
            worksheet[f"{col_letter}1"].fill = header_fill

    st.success(f"✅ {len(df_final)} questões geradas com sucesso!")

    st.download_button(
        label="📥 Baixar XLSX",
        data=output.getvalue(),
        file_name=f"{nome_arquivo}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

def gerar_csv_udemy(texto, nome_arquivo):
    output = io.StringIO()
    questions = []

    blocks = re.split(r"(?=Question \d+\n)", texto.strip())

    for block in blocks:
        lines = block.strip().splitlines()
        if not lines or not lines[0].startswith("Question"):
            continue

        question_text = ""
        answers = []
        correct_indexes = []
        in_question = False

        for i, line in enumerate(lines):
            line = line.strip()

            if re.match(r"^Question \d+", line):
                in_question = True
                continue

            if in_question and not question_text and line and line.upper() != "SKIPPED":
                question_text = line
                continue

            if "Correct answer" in line or "Correct selection" in line:
                if i + 1 < len(lines):
                    answer_text = lines[i + 1].strip()
                    if answer_text not in answers:
                        answers.append(answer_text)
                    correct_indexes.append(answers.index(answer_text) + 1)
            elif line and not line.startswith("Overall explanation") and not line.startswith("Skipped") and not any(kw in line for kw in ["Correct answer", "Correct selection"]):
                if line not in answers:
                    answers.append(line)

        question_type = "multi-select" if len(correct_indexes) > 1 else "multiple-choice"

        qdata = {
            "Question": question_text,
            "Question Type": question_type,
            "Correct Answers": ",".join(map(str, correct_indexes)),
            "Overall Explanation": "",
            "Domain": ""
        }

        for i in range(6):
            qdata[f"Answer Option {i+1}"] = answers[i] if i < len(answers) else ""
            qdata[f"Explanation {i+1}"] = ""

        questions.append(qdata)

    df_csv = pd.DataFrame(questions, columns=CSV_HEADER)

    df_csv.to_csv(output, index=False)

    st.success(f"✅ {len(df_csv)} questões processadas para CSV!")

    st.download_button(
        label="📥 Baixar CSV para Udemy",
        data=output.getvalue(),
        file_name=f"{nome_arquivo}.csv",
        mime="text/csv"
    )
